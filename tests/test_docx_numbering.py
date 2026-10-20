"""Word's automatic list numbering (extract._DocxNumbering). A list made
with Word's numbering button has no number in its text - Word draws "I." /
"(a)" / "1." from numbering.xml - and 6 of the 7 real Word papers use it:
"paper 6.docx" (Grade 13 Biosystem Technology) numbers every MCQ option
that way, and read without the labels 0 of its 50 MCQs had options."""

import io
import sys
import zipfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.intake import extract

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _lvl(ilvl, fmt, text, start=1):
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/></w:lvl>')


# 1: questions "1." with lettered parts "(a)" under them.
# 2: MCQ options "I." - "IV."; numId 3 shares it but restarts at I, the way
#    Word saves "restart numbering" for the next question's options.
# 4: a bullet list.
NUMBERING = (
    f'<w:numbering {W}>'
    f'<w:abstractNum w:abstractNumId="1">{_lvl(0, "decimal", "%1.")}{_lvl(1, "lowerLetter", "(%2)")}</w:abstractNum>'
    f'<w:abstractNum w:abstractNumId="2">{_lvl(0, "upperRoman", "%1.")}</w:abstractNum>'
    f'<w:abstractNum w:abstractNumId="3">{_lvl(0, "bullet", "•")}</w:abstractNum>'
    '<w:num w:numId="1"><w:abstractNumId w:val="1"/></w:num>'
    '<w:num w:numId="2"><w:abstractNumId w:val="2"/></w:num>'
    '<w:num w:numId="3"><w:abstractNumId w:val="2"/>'
    '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
    '<w:num w:numId="4"><w:abstractNumId w:val="3"/></w:num>'
    '</w:numbering>'
)


def _numbered(doc, text, num_id=None, ilvl=0):
    p = doc.add_paragraph(text)
    if num_id is not None:
        num_pr = OxmlElement("w:numPr")
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        num_pr.append(lvl)
        num_pr.append(nid)
        p._p.get_or_add_pPr().append(num_pr)
    return p


def _save(doc, path):
    """Save, then put our list definitions in as word/numbering.xml (the
    blank template python-docx starts from already has that part)."""
    buf = io.BytesIO()
    doc.save(buf)
    src = zipfile.ZipFile(buf)
    with zipfile.ZipFile(path, "w") as out:
        for item in src.infolist():
            data = NUMBERING.encode("utf-8") if item.filename == "word/numbering.xml" else src.read(item)
            out.writestr(item, data)
    return path


def _text(path):
    def no_ocr(*a, **k):
        raise extract.ExtractionUnavailable("no OCR in this test")
    text, _note, _bold = extract.read_docx(path, no_ocr)
    return [line for line in text.split("\n") if line.strip()]


def test_options_restart_for_every_question(tmp_path):
    d = docx.Document()
    d.add_paragraph("01. Water quality determination measures,")
    for o in ("Only temperature", "pH, nutrients and contaminants", "Air pressure", "Soil texture"):
        _numbered(d, o, num_id=2)
    d.add_paragraph("02. Plantation crop production techniques emphasize,")
    for o in ("Ignoring quality standards", "Efficient processing", "Reducing yield", "Avoiding markets"):
        _numbered(d, o, num_id=3)
    assert _text(_save(d, tmp_path / "mcq.docx")) == [
        "01. Water quality determination measures,",
        "I. Only temperature", "II. pH, nutrients and contaminants", "III. Air pressure", "IV. Soil texture",
        "02. Plantation crop production techniques emphasize,",
        "I. Ignoring quality standards", "II. Efficient processing", "III. Reducing yield", "IV. Avoiding markets",
    ]


def test_questions_and_their_lettered_parts(tmp_path):
    d = docx.Document()
    _numbered(d, "A particle is projected vertically upwards.", num_id=1)
    _numbered(d, "Find the greatest height reached.", num_id=1, ilvl=1)
    _numbered(d, "Find the time taken to return.", num_id=1, ilvl=1)
    d.add_paragraph("A note between the questions.")
    _numbered(d, "A projectile is fired at an angle.", num_id=1)
    _numbered(d, "Find its range.", num_id=1, ilvl=1)
    _numbered(d, "A list item.", num_id=4)
    assert _text(_save(d, tmp_path / "parts.docx")) == [
        "1. A particle is projected vertically upwards.",
        "(a) Find the greatest height reached.", "(b) Find the time taken to return.",
        "A note between the questions.",
        # The same list carries on after a plain paragraph, and the
        # letters start again under the next question.
        "2. A projectile is fired at an angle.", "(a) Find its range.",
        "A list item.",  # a bullet adds no label
    ]


def test_an_auto_numbered_mcq_gets_its_options(tmp_path):
    d = docx.Document()
    d.add_paragraph("Part I")
    d.add_paragraph("Choose the correct answer.")
    for n, (stem, num_id) in enumerate((("Water quality determination measures,", 2),
                                         ("Plantation crop production techniques emphasize,", 3)), 1):
        d.add_paragraph(f"0{n}. {stem}")
        for o in ("Only temperature", "pH, nutrients and contaminants", "Air pressure", "Soil texture"):
            _numbered(d, o, num_id=num_id)
    path = _save(d, tmp_path / "paper.docx")
    drafts = extract.extract(str(path), ocr=lambda *a: "", filename="paper.docx")["questions"]
    assert [len(extract.split_mcq_options(q["text"])[1]) for q in drafts] == [4, 4]


def test_a_file_without_list_definitions_reads_as_before(tmp_path):
    d = docx.Document()
    d.add_paragraph("1. Plain typed numbering.")
    buf = io.BytesIO()
    d.save(buf)
    src = zipfile.ZipFile(buf)
    path = tmp_path / "plain.docx"
    with zipfile.ZipFile(path, "w") as out:
        for item in src.infolist():
            if item.filename != "word/numbering.xml":
                out.writestr(item, src.read(item))
    assert extract._docx_numbering(path) is None
