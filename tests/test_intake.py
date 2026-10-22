"""Paper intake - extraction heuristics, the batch duplicate check, and
the intake storage.

No real model, no real PDF: the pure text rules (segment,
detect_paper_meta) are exercised directly, the batch check runs against
a fake embedder the same way test_check_new.py does, and read_docx runs
against a document python-docx builds in tmp_path.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import store
from src.intake import extract, review


# --------------------------------------------------------- segment()

def test_segment_splits_on_numbered_questions():
    text = (
        "1. What is the capital of Sri Lanka? (2 marks)\n"
        "2. Name the largest ocean on Earth. (3 marks)\n"
        "3. Explain the water cycle. (5 marks)\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "2", "3"]
    assert [d["marks"] for d in drafts] == [2, 3, 5]
    assert drafts[0]["seq"] == 1
    assert "capital of Sri Lanka" in drafts[0]["text"]


def test_segment_handles_q_prefix_and_parens():
    for text in (
        "Q1 Define photosynthesis. [4]\nQ2 State Newton's first law. [3]\n",
        "(1) Define photosynthesis. [4]\n(2) State Newton's first law. [3]\n",
    ):
        drafts = extract.segment(text)
        assert [d["question_number"] for d in drafts] == ["1", "2"]
        assert [d["marks"] for d in drafts] == [4, 3]


def test_segment_handles_question_word_before_a_parenthesised_number():
    """The real GIT A/L model paper: Part II numbers its first five
    questions as bare "(01)".."(05)", then switches to "Question (06)"
    and "Question (07)" for the last two - no Part heading in between,
    and all seven are genuinely real questions. Losing the last two to
    the "question word" not being recognised as the same bracketed
    numbering left them glued onto question 5's text."""
    text = (
        "(01) A coastal town installs a flood warning system.\n"
        "(02) Ravindu is a software developer.\n"
        "(03) Senuri is building a personal web page.\n"
        "(04) Pavan works for a national examination department.\n"
        "(05) Chalani is researching the impacts of ICT.\n"
        "Question (06) [15 marks]\n"
        "Amara is a software engineer building a banking platform.\n"
        "Question (07) [15 marks]\n"
        "Dilshan is developing a multimedia web portfolio.\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == [
        "01", "02", "03", "04", "05", "06", "07"]
    # The "[15 marks]" annotation is not stripped from the body - only the
    # leading number marker is (same as every other question here) - but
    # it must still be captured as this question's marks.
    assert "Amara is a software engineer" in drafts[5]["text"]
    assert "Dilshan is developing" in drafts[6]["text"]
    assert drafts[5]["marks"] == 15
    assert drafts[6]["marks"] == 15


def test_mcq_options_on_their_own_lines_stay_inside_the_question():
    """The real Grade 12 GIT paper: questions "1." with options "(1)" to
    "(4)" one per line. Any left-margin number counted as a question, so
    options (2), (3) and (4) of question 1 became questions 2, 3 and 4,
    and the real questions 2-4 were dropped as going backwards."""
    text = (
        "1. Which of the following BEST describes a system?\n"
        "(1) A device\n(2) An organized set\n(3) A group\n(4) Any machine\n"
        "2. Which correctly represents IPO?\n"
        "(1) Output\n(2) Input\n(3) Process\n(4) Store\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "2"]
    assert "(4) Any machine" in drafts[0]["text"]


def test_segment_ignores_a_number_that_jumps_backwards():
    """A '1990.' at a line start inside prose is not question 1990 - a
    real paper only counts up."""
    text = (
        "1. In which year did the event happen?\n"
        "1990. was a significant year for many reasons that we will not\n"
        "2. What caused it?\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "2"]
    assert "1990. was a significant year" in drafts[0]["text"]


def test_missing_marks_is_not_low_confidence():
    """A real MCQ paper prints no per-question marks at all. Flagging
    that put a low-confidence warning on all 50 questions of a real
    Grade 13 paper and told the reviewer nothing."""
    drafts = extract.segment("1. A question with no marks annotation at all.\n")
    assert drafts[0]["marks"] is None
    assert drafts[0]["extraction_confidence"] == "ok"


def test_almost_no_text_is_low_confidence():
    """Short text is the signal that the extractor actually struggled."""
    drafts = extract.segment("1. ?\n2. A properly worded question here.\n")
    assert drafts[0]["extraction_confidence"] == "low"
    assert drafts[1]["extraction_confidence"] == "ok"


def test_segment_handles_number_formats_pdf_text_actually_produces():
    """The strict first version only split '1. ' with a trailing space.
    Real extracted text uses all of these."""
    cases = [
        "1.What is X?\n2.What is Y?",                 # no space after the dot
        "Question 1: Define X.\nQuestion 2: Define Y.",  # word + colon
        "Q.1 Define X.\nQ.2 Define Y.",                # Q. prefix
        "01. Define X.\n02. Define Y.",                # leading zero
        "1.\nDefine X on the next line.\n2.\nDefine Y.",  # number alone on its line
    ]
    for text in cases:
        drafts = extract.segment(text)
        assert len(drafts) == 2, f"did not split: {text!r}"
        assert [d["question_number"] for d in drafts] == ["1", "2"] or \
               [d["question_number"] for d in drafts] == ["01", "02"]


def test_segment_does_not_split_on_mcq_option_numbers():
    """(1)(2)(3)(4) listed as answer options after a question must not
    each become a new question."""
    text = (
        "5) What is the capital of Sri Lanka?\n"
        "(1) Kandy (2) Colombo (3) Galle (4) Jaffna\n"
        "6) Name the largest ocean.\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["5", "6"]


def test_segment_does_not_read_a_mark_total_as_a_question():
    """'(20)' right before the word marks is a total, not question 20."""
    text = "1. Answer fully. (20) marks\n2. Next question. (10 marks)\n"
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "2"]


def test_segment_falls_back_to_one_draft_when_no_numbering():
    text = "This is just a blob of text with no question numbers anywhere."
    drafts = extract.segment(text)
    assert len(drafts) == 1
    assert drafts[0]["question_number"] is None
    assert drafts[0]["extraction_confidence"] == "low"
    assert drafts[0]["text"] == text


def test_segment_takes_the_last_marks_annotation():
    """A multi-part question lists part marks then a total - the total
    is the useful one."""
    text = "1. (a) Define X. (2 marks) (b) Explain Y. (3 marks) Total (5 marks)\n"
    drafts = extract.segment(text)
    assert drafts[0]["marks"] == 5
    # The annotation used to read that 5 is gone from the stored text -
    # a student's own answer will never contain "Total (5 marks)", so
    # leaving it in would corrupt a future comparison against one.
    assert "Total (5 marks)" not in drafts[0]["text"]
    assert "Define X." in drafts[0]["text"]


def test_bare_parenthesised_marks_at_the_end_of_a_line_are_read():
    """A real Grade 12 Sinhala-medium Business Statistics paper marks
    every part this way - no "marks" word at all (the paper is entirely
    in Sinhala) and no square brackets either, just "(4)" at the very
    end of the line."""
    text = "1. Define the term.    (4)\n2. Explain the process.    (6)\n"
    drafts = extract.segment(text)
    assert drafts[0]["marks"] == 4
    assert drafts[0]["text"] == "Define the term."
    assert drafts[1]["marks"] == 6


def test_a_bare_parenthesised_number_mid_sentence_is_not_read_as_marks():
    """The same shape the rule above reads as marks must stay alone when
    it sits mid-sentence - the "12 cm" / "in 1948" problem the original
    bare-number rule was written to avoid, just parenthesised this time."""
    text = "1. The committee of (12) members then reviewed the proposal.\n"
    drafts = extract.segment(text)
    assert drafts[0]["marks"] is None


def test_marks_annotation_is_stripped_from_a_leaf_questions_text():
    text = "1. Convert the number to binary. [3]\n2. Explain the result. [2]\n"
    drafts = extract.segment(text)
    assert drafts[0]["text"] == "Convert the number to binary."
    assert drafts[0]["marks"] == 3
    assert drafts[1]["text"] == "Explain the result."
    assert drafts[1]["marks"] == 2


def test_marks_annotation_is_stripped_from_a_split_parts_own_text():
    text = (
        "1. A farmer grows two crops on the same land.\n"
        "(i) Explain what is meant by intercropping in this context. [3]\n"
        "(ii) State two advantages of the practice for the farmer. [4]\n"
        "(iii) Describe one problem the farmer may face with it. [3]\n"
    )
    drafts = extract.segment(text)
    by_number = {d["question_number"]: d for d in drafts}
    assert by_number["1(i)"]["text"] == (
        "Explain what is meant by intercropping in this context.")
    assert by_number["1(i)"]["marks"] == 3
    assert "[3]" not in by_number["1(i)"]["text"]


def test_marks_annotation_is_stripped_from_the_answer_text():
    """The marking scheme's own "[N marks]" label must not survive into
    the stored expected answer, for the same reason as the question
    text - it is not part of what a student would ever write."""
    parsed = extract.structure(_TWO_PART_PAPER)
    q01 = _part(parsed, "Part II")[0]
    parts = {p["question_number"]: p for p in _under(parsed, q01)}
    answer = parts["01(1)(a)"]["answer"]
    assert answer.startswith("Definition of a System")
    assert "[2 marks]" not in answer
    assert "marks" not in answer.lower()


def test_segment_strips_the_printed_question_number_from_the_text():
    """The number is metadata (question_number), not part of the
    question - leaving it in would break exact-match de-duplication and
    misfire the numbers guard."""
    drafts = extract.segment(
        "3. Define osmosis. (2 marks)\n7. Define osmosis. (2 marks)\n"
    )
    assert drafts[0]["text"].startswith("Define osmosis")
    assert drafts[1]["text"].startswith("Define osmosis")
    # Same question under two different printed numbers -> identical text.
    assert drafts[0]["text"] == drafts[1]["text"]
    assert [d["question_number"] for d in drafts] == ["3", "7"]


def test_real_mcq_paper_shape_splits_on_every_question():
    """The shape of a real Grade 13 Biosystem Technology paper: '01.' with
    no space, options on the lines below. The first extractor required a
    space after the dot and returned the whole 18,000-character paper as
    one question."""
    text = (
        "BIOSYSTEM TECHNOLOGY ( I , II )\nThird term examination\nGrade 13\n"
        "\tPart I\n"
        "01.Interpretes weather conditions suitable for biological systems,\n"
        "Ignoring climate\nUsing forecasts\nBoth of these\n"
        "02.Water quality determination measures,\n"
        "Only temperature\npH and nutrients\nNeither\n"
        "03.Plantation crop production techniques emphasize,\n"
        "Ignoring quality standards\nMeeting them\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["01", "02", "03"]
    assert drafts[0]["text"].startswith("Interpretes weather")
    # The options belong to their question, not to a new one.
    assert "Ignoring climate" in drafts[0]["text"]


def test_worked_solution_numbering_is_not_split_by_the_question_word_style():
    """A marking-scheme document numbers its SOLUTION STEPS, which look
    exactly like question numbers. Measured on a real Grade 13 Combined
    Maths upload: number-splitting found only solution steps. Splitting
    on the word 'Question' finds the real questions instead."""
    text = (
        "Question\nA particle is projected at 40 m/s.\n"
        "1. Initial Velocity Components\nvx = 40 cos a\n"
        "2. Finding the Vertical Velocity\nvy = u sin a - gt\n"
        "3. Magnitude of Resultant Velocity\nV = sqrt(vx^2 + vy^2)\n"
        "Question\nA string passes over a smooth pulley.\n"
        "1. Resolving Forces\n2mg - T = 2ma\n"
        "2. Taking Moments\nT = m(g + a)\n"
    )
    by_number = extract.segment(text, style="number")
    by_word = extract.segment(text, style="question_word")

    # Number-splitting is fooled - it finds the solution steps.
    assert len(by_number) > 2
    # The word style finds the two real questions.
    assert len(by_word) == 2
    assert by_word[0]["text"].startswith("A particle is projected")
    assert by_word[1]["text"].startswith("A string passes over")


def test_split_styles_are_declared_for_the_ui():
    """The page builds its radio buttons from these, so an unlabelled or
    missing style would render an empty choice."""
    assert extract.DEFAULT_SPLIT_STYLE in extract.SPLIT_STYLES
    assert "question_word" in extract.SPLIT_STYLES
    assert all(isinstance(v, str) and v for v in extract.SPLIT_STYLES.values())


# ------------------------------------------------------- structure()

# The shape of the real Grade 12 GIT first-term paper, cut down: a header,
# Part I (MCQs "1." with options "(1)"-"(4)"), Part II (structured "(01)"
# with parts "(1)", "(a)"), then a marking scheme for both.
_TWO_PART_PAPER = (
    "General Information Technology (GIT)\n"
    "Grade 12 — First Term Examination\n"
    "Time: Three Hours\n"
    "General Information Technology I\n"
    "Instructions:\n"
    "* Answer all the questions.\n"
    "* In each of the questions 1 to 40, pick one of the alternatives (1), (2), (3), (4) which is correct or most\n"
    "appropriate and mark your response on the answer sheet.\n"
    "1. Which of the following BEST describes a \"system\"?\n"
    "(1) A device that only performs a single function\n"
    "(2) An organized set of interacting components\n"
    "(3) A group of unrelated hardware parts\n"
    "(4) Any machine that uses electricity\n"
    "2. Which of the following devices can function as BOTH an input and an output device?\n"
    "(1) Keyboard\n(2) Touchscreen\n(3) Scanner\n(4) Plotter\n"
    "General Information Technology II\n"
    "* Answer FOUR questions only.\n"
    "* Each question carries 15 marks.\n"
    "(01) (1) A school library uses a computerized system to manage its book lending.\n"
    "(a) Define the term \"system\" as used in computing. [2]\n"
    "(b) Draw a fully labeled block diagram of a computer system. [5]\n"
    "(2) Consider the following scenario:\n"
    "(a) Identify ONE example of data and ONE example of information. [2]\n"
    "(b) Name and briefly describe any FOUR attributes of quality "
    "information. [2]\n"
    "(02) (1) Nimal is studying number systems in his GIT class.\n"
    "(a) Convert the decimal number 53 into binary.\n"
    "[3]\n"
    "MARKING SCHEME\n"
    "Grade 12 — General Information Technology — First Term Examination\n"
    "Part I — Multiple Choice Question Answers\n"
    "Q1: (2) Q2: (2)\n"
    "Answer Justifications for Selected MCQ:\n"
    "Q2: Touchscreen both displays and accepts input → Answer (2)\n"
    "Part II — Marking Scheme (Each question: 15 marks)\n"
    "Question (01) [15 marks]\n"
    "(1)(a) Definition of a System [2 marks]\n"
    "A system is an organized collection of interacting components.\n"
    "Question (02) [15 marks]\n"
    "(1)(a) Convert Decimal 53 to Binary [3 marks]\n"
    "53₁₀ = 110101₂\n"
)


def _part(parsed, label):
    """A part's questions - its own, not the parts inside them, which are
    drafts in their own right now (see _split_parts)."""
    return [q for q in parsed["questions"]
            if q["section"] == label and not q["depth"]]


def _under(parsed, question):
    """The drafts nested under one question, in reading order."""
    seqs, out = {question["seq"]}, []
    for q in parsed["questions"]:
        if q["parent_seq"] in seqs:
            seqs.add(q["seq"])
            out.append(q)
    return out


def test_a_two_part_paper_comes_back_as_its_two_parts():
    """Read as one list, the real GIT paper came back as 54 "questions"
    for 45 real ones - Part II's parts split apart and the marking scheme
    split into questions of its own."""
    parsed = extract.structure(_TWO_PART_PAPER)
    assert [(s["label"], s["kind"], s["questions"])
            for s in parsed["sections"]] == [
        ("Part I", "mcq", 2), ("Part II", "structured", 2)]

    part1, part2 = _part(parsed, "Part I"), _part(parsed, "Part II")
    assert [q["question_number"] for q in part1] == ["1", "2"]
    assert "(4) Any machine" in part1[0]["text"]
    assert [q["question_number"] for q in part2] == ["01", "02"]
    # "Each question carries 15 marks", not the last part's "[2]".
    assert [q["marks"] for q in part2] == [15, 15]
    assert [q["seq"] for q in parsed["questions"]] == list(
        range(1, len(parsed["questions"]) + 1))


def test_a_structured_question_is_split_into_its_parts():
    """The real GIT question (01) is a stem, two numbered parts, and
    lettered parts under those - each with its own marks, and each the
    thing a teacher actually marks and answers. Read as one draft it was
    a single box with a single answer."""
    parsed = extract.structure(_TWO_PART_PAPER)
    q01 = _part(parsed, "Part II")[0]
    parts = _under(parsed, q01)

    assert [(p["question_number"], p["depth"]) for p in parts] == [
        ("01(1)", 1),
        ("01(1)(a)", 2),
        ("01(1)(b)", 2),
        ("01(2)", 1),
        ("01(2)(a)", 2),
        ("01(2)(b)", 2),
    ]
    # The stem stays with the part it introduces, not with its own parts.
    assert parts[0]["text"].startswith("A school library")
    assert parts[1]["text"].startswith("Define the term")
    # Every part points at the draft it sits under.
    by_number = {p["question_number"]: p for p in parts}
    assert by_number["01(1)(a)"]["parent_seq"] == by_number["01(1)"]["seq"]
    assert by_number["01(1)"]["parent_seq"] == q01["seq"]


def test_each_part_carries_its_own_marks():
    parsed = extract.structure(_TWO_PART_PAPER)
    parts = _under(parsed, _part(parsed, "Part II")[0])
    marks = {p["question_number"]: p["marks"] for p in parts}
    assert marks["01(1)(a)"] == 2
    assert marks["01(1)(b)"] == 5


def test_split_question_marks_sum_to_the_question_total():
    """The real GIT A/L paper prints no marks total for "(1)" or "(2)"
    themselves - only their own lettered parts carry a figure, e.g.
    (1)(a) [2], (1)(b) [5]. Before this, "(1)"'s own marks field showed
    5 - (b)'s own figure, the last one found searching (1)'s whole text
    - instead of 7, what (a) and (b) actually add up to."""
    parsed = extract.structure(_TWO_PART_PAPER)
    parts = _under(parsed, _part(parsed, "Part II")[0])
    marks = {p["question_number"]: p["marks"] for p in parts}
    assert marks["01(1)"] == 7   # (a) 2 + (b) 5
    assert marks["01(2)"] == 4   # (a) 2 + (b) 2


def test_split_question_marks_are_blank_when_a_part_has_none_yet():
    """If even one part's own mark is unknown, summing the rest would
    understate the real total and look no different from a genuine one
    - left blank instead, so a reviewer sees a gap to fill rather than
    a wrong number to trust."""
    text = (
        "1. A farmer grows two crops on the same land.\n"
        "(i) Explain what is meant by intercropping in this context. [3]\n"
        "(ii) State two advantages of the practice for the farmer.\n"
        "(iii) Describe one problem the farmer may face with it. [3]\n"
    )
    drafts = extract.segment(text)
    assert drafts[0]["marks"] is None


def test_mcq_options_are_never_split_into_parts():
    """Options are a numbered list at the left margin, exactly like a
    part list, but an option is not a question with an answer of its
    own."""
    parsed = extract.structure(_TWO_PART_PAPER)
    for q in _part(parsed, "Part I"):
        assert _under(parsed, q) == []
    assert "(4) Any machine" in _part(parsed, "Part I")[0]["text"]


def test_mcq_with_verbose_roman_numeral_options_is_not_split_into_parts():
    """The real Grade 13 Biosystem Technology paper numbers every Part I
    option "I. II. III. IV. V." rather than "(1) (2) (3)...". A question
    whose options happened to be short text stayed whole by luck (too
    short to look like a real structured part), but one with longer,
    descriptive options - like this real question 9 - was read as a
    structured question and split_parts()'d into five fake sub-
    questions, each stripped of its own "I."/"II." as if it were a real
    part's printed number rather than an MCQ option label."""
    text = (
        "Part I\n"
        "09.What is mean by egg's shape index?\n"
        "I. The percentage of ratio between the length and the breadth "
        "that's important in packaging\n"
        "II. The percentage of ratio between the length and the weight "
        "that's important in packaging\n"
        "III. The percentage of ratio between the length and the breadth "
        "that's important in deciding the quality while cooking\n"
        "IV. The percentage of ratio between the length and the weight "
        "that's important in deciding the quality while cooking\n"
        "V. The percentage of ration between the breadth and the weight "
        "that's important in deciding the brittleness\n"
        "10.Which of the following is an example for temporary controlled "
        "environmental structure?\n"
        "I. Web house\nII. Solar greenhouse\nIII. Green house\n"
        "IV. Polythene greenhouse\nV. Lath house\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["09", "10"]
    stem, options = extract.split_mcq_options(drafts[0]["text"])
    assert stem == "What is mean by egg's shape index?"
    assert [o["label"] for o in options] == ["I", "II", "III", "IV", "V"]


def test_nested_roman_sub_questions_are_not_misread_as_mcq_options():
    """Adding roman numerals to _OPTION_LINE (previous test) broke a
    DIFFERENT part of this same real paper: Part II-A nests roman-
    numeral sub-questions inside lettered parts - "a) i)Mention the
    main processes...", "b) i) Give examples...". "a)"/"b)" themselves
    started counting as MCQ option lines, because the verb that should
    have excluded them ("Mention", "Give") sits after ANOTHER marker
    ("i)"), not right after "a)"/"b)" the way the exclusion check
    expects - so the whole section's kind flipped to "mcq", which
    skips part-splitting entirely: question 04 never split into real
    a)/b)/i)/ii)/iii) parts at all, staying one undivided block."""
    text = (
        "Part II-A\n"
        "04.\n"
        "a) i)List down the disadvantages of frrezing.\n"
        "ii) What is the principle applies in blanching?\n"
        "iii) Give examples of food diversification.\n"
        "b) i) What are the reasons of food packaging?\n"
        "ii) List down the advantages of modified atmosphere packaging.\n"
        "iii) Writedown the measureable factors of post harvest losses.\n"
    )
    drafts = extract.segment(text)
    numbers = [d["question_number"] for d in drafts]
    assert numbers == [
        "04", "04(a)", "04(a)(i)", "04(a)(ii)", "04(a)(iii)",
        "04(b)", "04(b)(i)", "04(b)(ii)", "04(b)(iii)",
    ]
    by_number = {d["question_number"]: d for d in drafts}
    assert by_number["04(a)(i)"]["text"] == (
        "List down the disadvantages of frrezing.")
    assert by_number["04(b)(iii)"]["text"].startswith(
        "Writedown the measureable factors")


def test_split_mcq_options_separates_stem_from_each_option():
    """A review page wants each option as its own card - unlike
    structure() itself, which deliberately keeps an MCQ's stem and
    options as one piece of text for the duplicate-check fingerprint."""
    parsed = extract.structure(_TWO_PART_PAPER)
    q1 = _part(parsed, "Part I")[0]
    stem, options = extract.split_mcq_options(q1["text"])
    assert stem == 'Which of the following BEST describes a "system"?'
    assert [(o["label"], o["text"]) for o in options] == [
        ("1", "A device that only performs a single function"),
        ("2", "An organized set of interacting components"),
        ("3", "A group of unrelated hardware parts"),
        ("4", "Any machine that uses electricity"),
    ]


def test_split_mcq_options_handles_five_lettered_options():
    """Whatever the paper numbers them with, and however many there
    are - not hardcoded to four."""
    body = ("Which gas is most responsible for the greenhouse effect?\n"
            "A. Nitrogen\nB. Oxygen\nC. Carbon dioxide\nD. Argon\nE. Helium")
    stem, options = extract.split_mcq_options(body)
    assert stem == "Which gas is most responsible for the greenhouse effect?"
    assert [o["label"] for o in options] == ["A", "B", "C", "D", "E"]
    assert options[2]["text"] == "Carbon dioxide"


def test_split_mcq_options_ignores_a_single_stray_bracket():
    """One bracketed word is not an option list - the same "two markers
    minimum" rule _split_parts() already relies on."""
    body = "Define the term (system) as used in computing. [2]"
    stem, options = extract.split_mcq_options(body)
    assert stem == body
    assert options == []


def test_split_mcq_options_handles_capital_roman_numerals():
    """Whichever way a paper numbers its options - capital roman
    numerals included, bracketed or not."""
    body = ("Which of these is a noble gas?\n"
            "I. Oxygen\nII. Nitrogen\nIII. Helium\nIV. Chlorine")
    stem, options = extract.split_mcq_options(body)
    assert [o["label"] for o in options] == ["I", "II", "III", "IV"]
    assert options[2]["text"] == "Helium"

    bracketed = ("Which pathway does NOT lead to ATP production?\n"
                "(I) Glycolysis\n(II) Krebs cycle\n"
                "(III) Photosynthesis\n(IV) Electron transport chain")
    _, options2 = extract.split_mcq_options(bracketed)
    assert [o["label"] for o in options2] == ["I", "II", "III", "IV"]


def test_split_mcq_options_keeps_lettered_statements_in_the_stem():
    """A stem can hold its own lettered list of statements before a
    Roman-numeral option list. The statements are part of the question,
    not options - the options are the list that comes last."""
    body = ("Identify the correct combination of associated factors "
            "related to ornamental fish breeding.\n"
            "a. Platty and Swordtail are live bearing fish, during breeding "
            "they give birth to fry.\n"
            "b. In ornamental fish breeding, it is necessary to maintain a "
            "suitable environment with proper water quality.\n"
            "c. When transporting ornamental fish in a plastic bag, the "
            "water should be enriched with oxygen.\n"
            "d. When ornamental fish are injured, fungal infections can be "
            "observed in wounds.\n"
            "The correct combination of associated factors is,\n"
            "I. a,c only\nII. c,d only\nIII. a,b only\nIV. b,c only\n"
            "V. a,d only")
    stem, options = extract.split_mcq_options(body)
    assert [o["label"] for o in options] == ["I", "II", "III", "IV", "V"]
    assert options[2]["text"] == "a,b only"
    assert stem.startswith("Identify the correct combination")
    assert "d. When ornamental fish" in stem


def test_bare_digit_options_are_found_as_a_last_resort():
    """A real Grade 12 Sinhala-medium Business Statistics paper's MCQs
    number every option "1. 2. 3. 4. 5." with no brackets at all - the
    one shape every earlier style in split_mcq_options() deliberately
    excludes (digits only ever count bracketed, to protect a
    worked-solution's own numbered SOLUTION STEPS from being mistaken
    for options or parts). Trusted here only because there are at least
    four of them."""
    body = (
        "Which gas is most responsible for the greenhouse effect?\n"
        "1. Nitrogen\n2. Oxygen\n3. Carbon dioxide\n4. Argon\n5. Helium"
    )
    stem, options = extract.split_mcq_options(body)
    assert stem == "Which gas is most responsible for the greenhouse effect?"
    assert [o["label"] for o in options] == ["1", "2", "3", "4", "5"]
    assert options[2]["text"] == "Carbon dioxide"


def test_bare_digit_options_need_at_least_four_to_count():
    """Fewer than four bare digits is exactly the numbered-solution-steps
    shape this rule exists to leave alone - a real worked-solutions
    document (Combined Maths) has just two steps per question."""
    body = (
        "A ball is thrown at 20 m/s at an angle of 30 degrees.\n"
        "1. Initial velocity components are found by resolving the vector.\n"
        "2. Finding the vertical velocity at the highest point of flight."
    )
    stem, options = extract.split_mcq_options(body)
    assert options == []
    assert stem == body


def test_bare_digit_options_do_not_swallow_a_preceding_lettered_statement_list():
    """The same real Sinhala paper's "combination" questions: lettered
    statements (A. B. C.) to judge, then bare-digit options combining
    them ("1. A only  2. A and B only..."). The statements must stay in
    the stem, not be mistaken for the options themselves - the same
    "lettered statements before the real option list" shape
    test_split_mcq_options_keeps_lettered_statements_in_the_stem already
    covers for a bracketed/Roman key, just with neither here."""
    body = (
        "Choose the answer containing only the true statement(s).\n"
        "A. Statement one is about data collection.\n"
        "B. Statement two is about business statistics.\n"
        "C. Statement three is about inferential statistics.\n"
        "\n"
        "1. A only     2. A and B only     3. B and C only     "
        "4. A and C only     5. A, B and C"
    )
    stem, options = extract.split_mcq_options(body)
    assert [o["label"] for o in options] == ["1", "2", "3", "4", "5"]
    assert options[0]["text"] == "A only"
    assert stem.startswith("Choose the answer")
    assert "C. Statement three" in stem


def test_sinhala_choose_the_answer_phrase_is_an_mcq_hint():
    """The real Grade 12 Sinhala-medium Business Statistics paper's own
    instructions ("...correct or best-matching answer, choose it") - the
    paper is entirely in Sinhala, so none of the English MCQ-hint phrases
    could ever match it otherwise, and without this the paper's own
    "A. B. C." lettered statements (a question of this exact
    "combination" shape) were wrongly split into fake sub-questions,
    since the section's kind could not be guessed any other way."""
    assert extract._MCQ_HINT.search(
        "1 සිට 50 තෙක් එක් එක් ප්‍රශ්නයට නිවැරදි හෝ ඉතාමත් ගැලපෙන "
        "පිළිතුර තෝරන්න."
    )


def test_mcq_correct_label_matches_every_real_label_style():
    assert extract.mcq_correct_label("(2)") == "2"
    assert extract.mcq_correct_label("A") == "a"
    # Roman numerals beyond two characters - "iii", "viii" - are real
    # option labels and must not be silently rejected as too long.
    assert extract.mcq_correct_label("(III)") == "iii"
    assert extract.mcq_correct_label("VIII") == "viii"
    assert extract.mcq_correct_label(None) is None
    assert extract.mcq_correct_label("") is None
    # Free text a data-entry person typed by hand is not a label at all.
    assert extract.mcq_correct_label("The control unit") is None


def test_roman_numbered_parts_are_split():
    text = (
        "1. A farmer grows two crops on the same land.\n"
        "(i) Explain what is meant by intercropping in this context. [3]\n"
        "(ii) State two advantages of the practice for the farmer. [4]\n"
        "(iii) Describe one problem the farmer may face with it. [3]\n"
    )
    drafts = extract.segment(text)
    # The question's own mark is what its three parts add up to (3+4+3),
    # not whichever part's own figure happens to sit last in its raw
    # text - see _own_marks() and test_split_question_marks_sum_to_the_
    # question_total below.
    assert [(d["question_number"], d["depth"], d["marks"]) for d in drafts] == [
        ("1", 0, 10),
        ("1(i)", 1, 3),
        ("1(ii)", 1, 4),
        ("1(iii)", 1, 3),
    ]
    assert drafts[0]["text"] == "A farmer grows two crops on the same land."


def test_lettered_parts_without_brackets_are_split():
    text = (
        "1. Study the diagram of the water cycle shown above.\n"
        "a. Name the process that turns water into vapour and say where it "
        "happens. [2]\n"
        "b. Explain why the water returns to the ground as rain. [3]\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "1(a)", "1(b)"]


def test_sinhala_vowel_lettered_parts_are_split():
    """A real Grade 12 Sinhala-medium Business Statistics paper letters a
    structured question's parts with Sinhala's own independent vowels -
    "අ) ... ආ) ..." - the same role "(a) (b)" plays on an English paper.
    Roman-numeral sub-parts underneath stay in Latin script, unchanged."""
    text = (
        "01)\n"
        "අ)\n"
        "(i) Explain the first idea.    (4)\n"
        "(ii) Explain the second idea.    (4)\n"
        "\n"
        "ආ) A shared stem sentence for the next part.\n"
        "(i) Explain the third idea.    (4)\n"
        "(ii) Explain the fourth idea.    (4)\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == [
        "01", "01(අ)", "01(අ)(i)", "01(අ)(ii)",
        "01(ආ)", "01(ආ)(i)", "01(ආ)(ii)",
    ]
    by_number = {d["question_number"]: d for d in drafts}
    assert by_number["01"]["marks"] == 16
    assert by_number["01(ආ)"]["text"] == "A shared stem sentence for the next part."


def test_sinhala_vowel_lettering_can_skip_a_vowel():
    """The same real paper's question 2 letters its THIRD part "ඉ)",
    skipping two vowels in between ("ඇ)" and "ඈ)" never appear) -
    whether that is this paper's own inconsistency or a real, narrower
    convention is not something to assume from one paper, so the vowel
    sequence only has to move forward, not land on the very next vowel."""
    text = (
        "02)\n"
        "අ) First part stem.\n"
        "(i) Explain this.    (4)\n"
        "(ii) Explain this too.    (4)\n"
        "\n"
        "ආ) Second part stem.\n"
        "(i) Explain that.    (4)\n"
        "(ii) Explain that too.    (4)\n"
        "\n"
        "ඉ)\n"
        "(i) Explain the other thing.    (2)\n"
        "(ii) Explain one more thing.    (2)\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == [
        "02", "02(අ)", "02(අ)(i)", "02(අ)(ii)",
        "02(ආ)", "02(ආ)(i)", "02(ආ)(ii)",
        "02(ඉ)", "02(ඉ)(i)", "02(ඉ)(ii)",
    ]


def test_a_single_marker_does_not_split_a_question():
    """One "(a)" inside a sentence is not a part list."""
    text = "1. In figure (a) above, state the name of the labelled part. [2]\n"
    drafts = extract.segment(text)
    assert len(drafts) == 1
    assert drafts[0]["depth"] == 0


def test_numbered_solution_steps_are_not_parts():
    """A bare "1." at the margin is the paper's own numbering - on a
    worked-solutions document, the solution steps. Only bracketed digits
    count as parts, as the real GIT paper prints them."""
    text = (
        "Question 1\n"
        "A ball is thrown at 20 m/s at an angle of 30 degrees.\n"
        "1. Initial velocity components are found by resolving the vector.\n"
        "2. Finding the vertical velocity at the highest point of flight.\n"
    )
    drafts = extract.segment(text, style="question_word")
    assert len(drafts) == 1
    assert "Initial velocity components" in drafts[0]["text"]


def test_the_marking_scheme_becomes_answers_not_questions():
    parsed = extract.structure(_TWO_PART_PAPER)
    everything = " ".join(q["text"] for q in parsed["questions"])
    assert "MARKING SCHEME" not in everything
    assert "Q1: (2)" not in everything

    part1, part2 = _part(parsed, "Part I"), _part(parsed, "Part II")
    assert [q["answer"] for q in part1] == ["(2)", "(2)"]
    # The scheme marks a structured question part by part, and each part
    # is now its own draft with its own answer box - see the next test.
    by_number = {q["question_number"]: q for q in _under(parsed, part2[0])}
    assert by_number["01(1)(a)"]["answer"].startswith(
        "Definition of a System")
    # Question 02 has one marker at each level, so it is not split at all
    # (one marker is not a list) and keeps its whole worked answer.
    assert _under(parsed, part2[1]) == []
    assert "110101" in part2[1]["answer"]


def test_a_worked_answer_is_split_across_the_parts_it_marks():
    """"Question (01) ... (1)(a) Definition of a System [2 marks] ..." -
    the scheme labels each part's answer the way the paper labels the
    part, so it lands in that part's box rather than in one wall of text
    on the question."""
    parsed = extract.structure(_TWO_PART_PAPER)
    q01 = _part(parsed, "Part II")[0]
    parts = {p["question_number"]: p for p in _under(parsed, q01)}

    assert q01["answer"] is None          # nothing before the first label
    # "[2 marks]" is read into this part's own marks elsewhere - kept out
    # of the stored answer text itself, see test_marks_annotation_is_
    # stripped_from_the_answer_text.
    assert parts["01(1)(a)"]["answer"] == (
        "Definition of a System\n"
        "A system is an organized collection of interacting components.")
    # A part the scheme says nothing about keeps an empty box.
    assert parts["01(1)(b)"]["answer"] is None


def test_a_worked_answer_with_no_part_labels_stays_on_the_question():
    paper = (
        "Structured Questions\n"
        "(01) Explain how a computer stores a character in memory.\n"
        "(02) Describe two uses of a spreadsheet in a school office.\n"
        "MARKING SCHEME\n"
        "Question (01) [10 marks]\n"
        "Each character is stored as a binary code such as ASCII.\n"
        "Question (02) [10 marks]\n"
        "Recording marks, and preparing the timetable.\n"
    )
    drafts = extract.segment(paper)
    assert [d["answer"] for d in drafts] == [
        "Each character is stored as a binary code such as ASCII.",
        "Recording marks, and preparing the timetable.",
    ]


def test_answers_attach_when_the_marking_scheme_drops_the_parens():
    """The real GIT A/L paper: questions 6 and 7 are printed "Question
    (06)"/"Question (07)" (see the mixed-numbering test above), but its
    marking scheme heads their answers "Question 6"/"Question 7" - no
    leading zero, no parens at all. _attach_answers must still find them,
    since it compares question numbers as integers rather than strings."""
    paper = (
        "(01) A coastal town installs a flood warning system.\n"
        "Question (06) [15 marks]\n"
        "Amara is a software engineer building a banking platform.\n"
        "MARKING SCHEME\n"
        "Question (01) [15 marks]\n"
        "Rainfall sensors, timestamps.\n"
        "Question 6\n"
        "Control Unit fetches, ALU compares.\n"
    )
    drafts = extract.segment(paper)
    assert [d["question_number"] for d in drafts] == ["01", "06"]
    assert drafts[0]["answer"] == "Rainfall sensors, timestamps."
    assert drafts[1]["answer"] == "Control Unit fetches, ALU compares."


def test_a_parts_heading_and_instructions_stay_with_it():
    part2 = extract.structure(_TWO_PART_PAPER)["sections"][1]
    assert part2["heading"] == "General Information Technology II"
    assert part2["marks_each"] == 15
    assert "Answer FOUR questions only" in part2["instructions"]


def test_questions_before_a_part_2_heading_are_part_1():
    """The real Combined Maths paper has no heading over its first part,
    only "Part 2" before the second."""
    text = (
        "Combined mathematics- 2\nGrade -13\n1st term exam\n"
        "Answer the all questions\n"
        "1. Find x. [25 Marks]\n2. Find y. [25 Marks]\n"
        "Part 2\nAnswer any 5 questions\n1. Prove it. [25 Marks]\n"
    )
    parsed = extract.structure(text)
    assert [s["label"] for s in parsed["sections"]] == ["Part 1", "Part 2"]
    assert [q["section"] for q in parsed["questions"]] == [
        "Part 1", "Part 1", "Part 2"]


def test_a_paper_in_one_piece_has_no_part_labels():
    parsed = extract.structure("1. Define X.\n2. Define Y.\n")
    assert [s["label"] for s in parsed["sections"]] == [None]
    assert all(q["section"] is None for q in parsed["questions"])


def test_a_part_title_repeated_on_every_page_is_still_one_part():
    parsed = extract.structure(
        "Part I\n1. First question here.\nPart I\n2. Second question here.\n")
    assert [s["label"] for s in parsed["sections"]] == ["Part I"]
    assert [q["question_number"] for q in parsed["questions"]] == ["1", "2"]


def test_a_mid_part_instruction_moves_to_the_question_it_applies_to():
    """A real Grade 13 (A/L) Political Science paper has a bulleted
    instruction - "For questions 21 to 25, choose..." - sitting between
    question 20 and question 21, mid Part I, no heading involved.
    _number_starts() hands everything up to the next marker to the
    question before it, by construction, which wrongly glued this onto
    question 20's own text. It governs 21 through 25, so it becomes the
    instruction for that group - shown above question 21 - rather than
    part of any one question's text (see _split_group_instruction)."""
    text = (
        "20. A question.\n(1) One\n(2) Two\n\n"
        "• For questions 21 to 25, choose the correct combination.\n\n"
        "21. Another question.\n(1) ABCD\n(2) ACDE\n"
    )
    parsed = extract.structure(text)
    by_number = {q["question_number"]: q["text"] for q in parsed["questions"]}
    assert "For questions 21 to 25" not in by_number["20"]
    assert by_number["20"].strip().endswith("(2) Two")
    assert by_number["21"].startswith("Another question")
    assert [(g["from"], g["to"], g["text"]) for g in parsed["instruction_groups"]] == [
        ("21", "25", "• For questions 21 to 25, choose the correct combination.")]


def test_a_garbled_bullet_is_still_recognised_as_an_instruction():
    """OCR does not always read a bullet as itself - measured on the
    same real paper, Tesseract read it as a single stray Sinhala
    dependent vowel sign ("ේ") that cannot start a word on its own. Any
    Unicode combining mark works the same way, not just this one."""
    text = (
        "20. A question.\n(1) One\n(2) Two\n\n"
        "ෙ For questions 21 to 25, choose the correct combination.\n\n"
        "21. Another question.\n(1) ABCD\n(2) ACDE\n"
    )
    parsed = extract.structure(text)
    by_number = {q["question_number"]: q["text"] for q in parsed["questions"]}
    assert "For questions 21 to 25" not in by_number["20"]
    assert [(g["from"], g["to"]) for g in parsed["instruction_groups"]] == [("21", "25")]
    assert parsed["instruction_groups"][0]["text"].startswith("ෙ For questions 21 to 25")


def test_a_single_question_instruction_still_moves_to_that_question():
    """The older mover still handles an instruction that names no range -
    it belongs to the one question after it, so it goes to that question's
    front, not into a group."""
    text = (
        "20. A question.\n(1) One\n(2) Two\n\n"
        "• Read the passage below carefully before answering.\n\n"
        "21. Another question.\n(1) ABCD\n(2) ACDE\n"
    )
    parsed = extract.structure(text)
    by_number = {q["question_number"]: q["text"] for q in parsed["questions"]}
    assert parsed["instruction_groups"] == []
    assert by_number["21"].startswith("• Read the passage below")


def test_a_normal_short_last_paragraph_is_not_mistaken_for_an_instruction():
    """A question whose text genuinely ends in a short final paragraph -
    no bullet, no stray combining mark - must not be cut apart."""
    text = (
        "20. A question.\n(1) One\n(2) Two\n\nA final clarifying note.\n\n"
        "21. Another question.\n(1) ABCD\n(2) ACDE\n"
    )
    parsed = extract.structure(text)
    by_number = {q["question_number"]: q["text"] for q in parsed["questions"]}
    assert "A final clarifying note." in by_number["20"]
    assert "A final clarifying note." not in by_number["21"]


# --------------------------------------------- fill-in-the-blank

def test_split_fill_blanks_finds_numbered_inline_gaps():
    """A real Grade 12 General English paper: a whole paragraph with a
    preposition missing at five numbered points."""
    text = (
        "He went to the complex to prepare (1) ______ a tournament. He walked "
        "(2) ______ the field, then had a word (3) ______ his coach. "
        "(4) ______ the session he focused hard, and went (5) ______ the room."
    )
    stem, blanks = extract.split_fill_blanks(text)
    assert stem == text
    assert [b["number"] for b in blanks] == ["1", "2", "3", "4", "5"]


def test_split_fill_blanks_labels_by_the_enclosing_sentences_number():
    """A different shape, same paper: one blank per numbered sentence,
    with brackets of choices rather than a number right next to the gap
    itself - the label comes from whatever line the blank sits on."""
    text = (
        "1. The latest discoveries ______ (uncovered / revealed) evidence.\n"
        "2. Modern engineers design systems that are ______ (integrated / coordinated).\n"
    )
    stem, blanks = extract.split_fill_blanks(text)
    assert [b["number"] for b in blanks] == ["1", "2"]


def test_split_fill_blanks_needs_at_least_two():
    stem, blanks = extract.split_fill_blanks("Define X. ______")
    assert blanks == []


def test_split_fill_blanks_ignores_long_writing_space_lines():
    """A real Grade 12 General English paper's own open-ended writing
    questions use the same dots, just for a whole sentence or more, not
    one word - "(1) ....................................................
    .............................................................." runs
    over 100 characters. Must not be mistaken for a short blank."""
    text = (
        "(1) " + "." * 110 + "\n\n(2) " + "." * 110 + "\n\n(3) " + "." * 110
    )
    stem, blanks = extract.split_fill_blanks(text)
    assert blanks == []


def test_split_fill_blanks_ignores_dots_with_no_real_word_nearby():
    """A real Grade 12 General English paper mixes literal periods into
    an otherwise-ellipsis writing-space line ("...…"), which can be
    short enough on its own to look like a real blank by length alone -
    what gives it away is that nothing but more dots and blank lines
    sit beside it, never an actual word."""
    text = "…………………" + "." * 4 + "…\n\n…………" + "." * 5 + "…………\n\n" + "…" * 20
    stem, blanks = extract.split_fill_blanks(text)
    assert blanks == []


def test_fill_blank_answers_parses_bare_numbered_list():
    answer = "1. for \n2. across \n3. with \n4. During \n5. Into "
    assert extract.fill_blank_answers(answer) == {
        "1": "for", "2": "across", "3": "with", "4": "During", "5": "Into",
    }


def test_fill_blank_answers_parses_parenthesised_list():
    answer = "(1) documentation \n(2) original beginning \n(3) surprise"
    assert extract.fill_blank_answers(answer) == {
        "1": "documentation", "2": "original beginning", "3": "surprise",
    }


def test_fill_blank_answers_empty_for_plain_text():
    assert extract.fill_blank_answers("Paris") == {}
    assert extract.fill_blank_answers(None) == {}


def test_inline_blank_markers_are_not_read_as_mcq_options():
    """A real Grade 12 General English paper's fill-in-the-blank
    questions ("(1) ______ an important tournament... (2) ______ the
    entire field...") look exactly like the inline marker shape
    split_mcq_options()'s own fallback exists to catch for a real
    Tamil MCQ's options on one line - the two must not collide."""
    text = (
        "He went to prepare (1) ______ a tournament. He walked "
        "(2) ______ the field, then had a word (3) ______ his coach."
    )
    stem, options = extract.split_mcq_options(text)
    assert options == []
    assert extract.detect_question_type(text) == "fill_blank"


def test_a_blank_run_anywhere_in_an_option_disqualifies_it_too():
    """Not just right after the marker - a real Grade 12 General English
    paper's "respond to each situation" question has real prose first,
    then its own blank run further into the same numbered item (a
    separate one-sentence-response question, not a choice among
    options)."""
    text = (
        '(1) Your classmate failed an exam. Suggest a way to stay motivated.\n'
        '"Perhaps you ' + "." * 40 + '"\n'
        '(2) You won an award. A journalist asks about your feelings.\n'
        '"I ' + "." * 40 + '"\n'
        '(3) Someone offers you a dessert, but you are on a diet. Refuse politely.\n'
        '"No, ' + "." * 40 + '"\n'
    )
    stem, options = extract.split_mcq_options(text)
    assert options == []


def test_an_internally_blank_line_separated_option_disqualifies_it():
    """No blank run at all this time - a real Grade 12 General English
    vocabulary-in-context question numbers each word, but stacks its
    (unlabelled) 4-5 choice words with a blank line between them, inside
    what the marker shape alone makes look like one option. A genuine
    option is written as continuous text; the blank line always falls
    BETWEEN options, never inside one."""
    text = (
        "(1) reports\n\nspeech\noration\nstory\n\ndocumentation\n"
        "(2) origin\n\nnew ending\noriginal beginning\nfinal stage\nconclusion\n"
        "(3) wonder\n\nwalking\nsurprise\npanic\nfear\n"
    )
    stem, options = extract.split_mcq_options(text)
    assert options == []


def test_marking_scheme_with_no_label_is_still_found_by_shape():
    """A real Grade 12 General English paper's marking scheme carries no
    "Answers:"/"Marking Scheme:" label of its own at all - it just
    restates the exam's own "Part I" heading, then "Question 1" followed
    immediately by a flat numbered answer list, which a real question
    never is (real question text always follows "Question 1" instead)."""
    text = (
        "Part I\n"
        "Question 1\n"
        "Fill in the blanks using the words given. (05 marks)\n"
        "A student (1) ______ to the shop. He (2) ______ milk.\n"
        "Question 2\n"
        "Fill in each blank with the correct verb form. (05 marks)\n"
        "He (1) ______ (walk) home. He (2) ______ (buy) bread.\n"
        "\n\n\n"
        "Part I\n"
        "Question 1\n"
        "1. went\n2. bought\n"
        "Question 2\n"
        "1. walked\n2. bought\n"
    )
    parsed = extract.structure(text, style="question_word")
    by_number = {q["question_number"]: q for q in parsed["questions"] if not q.get("parent_seq")}
    assert len(by_number) == 2  # the scheme's own two "Question 1/2" did not become more questions
    assert extract.fill_blank_answers(by_number["1"]["answer"]) == {"1": "went", "2": "bought"}
    assert extract.fill_blank_answers(by_number["2"]["answer"]) == {"1": "walked", "2": "bought"}


def test_question_word_style_ignores_a_line_wrapped_mid_sentence_question():
    """A real Grade 12 General English paper: "...two sentences /
    \\nquestion in each blank." - a line wrap happens to land the lower-
    case word "question" at the very start of a line, mid-sentence. Only
    a genuine heading - "question" alone on its line, or followed by a
    real number - should split; prose continuing on the same line must
    not."""
    text = (
        "Question 7\n"
        "Complete the dialogue. Do not use more than two sentences /\n"
        "question in each blank. (05 marks)\n"
        "Ravi: Hello. \nSara: ......\n"
        "Question 8\n"
        "Write a paragraph.\n"
    )
    drafts = extract.segment(text, style="question_word")
    assert [d["question_number"] for d in drafts] == ["7", "8"]
    assert "in each blank" in drafts[0]["text"]


def test_lettered_sub_parts_are_not_mistaken_for_mcq_options():
    """The real Grade 13 Biosystem Technology paper's Part II-B: each
    question is "a) Explain ... b) Describe ...". Counted as options, the
    whole part was labelled multiple choice."""
    text = (
        "Part II – B\nAnswer any four questions.\n"
        "1. Irrigation\na) Explain the limitations of drip irrigation.\n"
        "b) Describe the method of sprinkler irrigation.\n"
        "c) State two advantages of furrow irrigation.\n"
        "2. Fish processing\na) Explain the steps in preparing jar fish.\n"
        "b) Compare salting and drying.\nc) Name two preservatives.\n"
    )
    assert extract.structure(text)["sections"][0]["kind"] != "mcq"


def test_underline_keeps_a_sentences_lettered_choices_as_mcq_options():
    """The real Grade 12 General English paper's question 6: "Underline
    the option that best expresses the meaning of each of the sentences
    given below" then five sentences, each followed by four lettered
    paraphrases to pick from. Read as parts, those four lettered items
    became four separate one-line "questions" each wanting its own
    answer box - wrong, since a student picks ONE, same as any MCQ.
    The word "underline" is what tells _split_parts() these lettered
    lists are choices, not further parts (see _UNDERLINE_WORD)."""
    text = (
        "1. Fill in the blank.\n2. Fill in another blank.\n"
        "3. Fill in a third blank.\n4. Fill in a fourth blank.\n"
        "5. Fill in a fifth blank.\n"
        "6. Underline the option that best expresses the meaning of each "
        "of the sentences given below in bold print.\n"
        "(1) Artificial intelligence is transforming the way businesses "
        "operate worldwide.\n"
        "a. AI has no impact on business.\n"
        "b. AI is slowly changing business methods.\n"
        "c. AI is significantly changing how businesses function globally.\n"
        "d. AI only affects technology companies.\n"
        "(2) The treaty was signed to prevent further environmental "
        "damage.\n"
        "a. The treaty caused environmental harm.\n"
        "b. The treaty was created to protect the environment.\n"
        "c. The treaty had no effect.\n"
        "d. The treaty encouraged pollution.\n"
    )
    drafts = extract.segment(text)
    by_number = {d["question_number"]: d for d in drafts}
    assert by_number["6(1)"]["text"].startswith(
        "Artificial intelligence is transforming")
    assert "a. AI has no impact" in by_number["6(1)"]["text"]
    # No fake "6(1)(a)" sub-question - the lettered list stayed merged.
    assert "6(1)(a)" not in by_number

    stem, options = extract.split_mcq_options(by_number["6(1)"]["text"])
    assert stem == "Artificial intelligence is transforming the way businesses operate worldwide."
    assert [o["label"] for o in options] == ["a", "b", "c", "d"]


def test_underline_does_not_swallow_a_genuinely_marked_part():
    """The hint word alone is not enough - a part that carries its own
    marks, or a further nested list, is still a real part, not an
    option, whatever the question's instruction says."""
    text = (
        "6. Underline the correct answer.\n"
        "a) Explain the main idea. [3]\n"
        "b) Explain the supporting detail. [4]\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["6", "6(a)", "6(b)"]


def test_a_leaf_part_is_detected_by_its_own_shape_not_its_depth():
    """A part one level deep with no children of its own - question 6's
    "6(1)" above - is exactly as answerable as a depth-0 question; only
    a part that itself HAS parts is a container ('sub'). Depth alone no
    longer forces 'sub' (see detect_question_type)."""
    mcq_text = (
        "Artificial intelligence is transforming the way businesses "
        "operate worldwide.\na. AI has no impact on business.\n"
        "b. AI is slowly changing business methods.\n"
        "c. AI is significantly changing how businesses function "
        "globally.\nd. AI only affects technology companies."
    )
    assert extract.detect_question_type(mcq_text, has_parts=False) == "mcq"
    # A container - a stem with real children - stays 'sub' regardless.
    assert extract.detect_question_type(mcq_text, has_parts=True) == "sub"


def test_a_single_stray_number_does_not_beat_a_real_numbering():
    """Worked solutions numbered like the questions they sit in restart
    over and over; that must not hand the paper to a format that numbers
    one thing once."""
    text = (
        "1. Find the velocity.\n1. Resolve.\n2. Solve.\n3. Conclude.\n"
        "2. Find the force.\n1. Resolve.\n2. Solve.\n"
        "3. Find the time.\n(4) An equation label on its own line\n"
    )
    numbers = [q["question_number"] for q in extract.segment(text)]
    assert len(numbers) >= 2
    assert "4" not in numbers


def test_heading_recognition():
    h = extract._heading
    assert h("General Information Technology II")["label"] == "Part II"
    assert h("\tPart I")["label"] == "Part I"
    assert h("Part II – B")["label"] == "Part II-B"
    assert h("Section A (Answer all questions)")["label"] == "Section A"
    assert h("MARKING SCHEME")["kind"] == "answers"
    # Lines from real papers that must not start a part.
    assert h("Part 1 of the experiment uses a convex lens.") is None
    assert h("Paper – 1 | Part 2 Questions with Answers") is None
    assert h("Answer") is None
    assert h("Answer the all questions") is None
    assert h("Explain the causes of World War II") is None


# ------------------------------------------------- detect_paper_meta()

def test_detect_paper_meta_reads_grade_subject_and_title():
    text = "Provincial Department of Education\nGrade 10 - Science\nFirst Term Test\n"
    meta = extract.detect_paper_meta(text)
    assert meta == {"grade": "G10", "subject": "Science",
                    "title": "First Term Test"}


def test_detect_paper_meta_prefers_the_longer_subject_name():
    text = "Grade 11\nBusiness Studies\nSecond Term\n"
    meta = extract.detect_paper_meta(text)
    assert meta["subject"] == "Business Studies"


def test_detect_paper_meta_returns_none_when_absent():
    meta = extract.detect_paper_meta("Just some text with no header at all.")
    assert meta == {"grade": None, "subject": None, "title": None}


def test_paper_meta_uses_the_banks_own_grade_and_subject_names():
    """The bank stores grades as G06-G13 and subjects under its own
    names. A detected "13" or "Combined Mathematics" matched nothing in
    the review page's lists, so grade and subject showed up blank."""
    maths = extract.detect_paper_meta(
        "Combined mathematics- 2\nGrade -13\n1st term exam\n")
    assert maths == {"grade": "G13", "subject": "Combined Maths",
                     "title": "1st Term Exam"}
    git = extract.detect_paper_meta(
        "General Information Technology (GIT)\n"
        "Grade 12 — First Term Examination\n")
    assert git == {"grade": "G12", "subject": "General Information Technology",
                   "title": "First Term Examination"}


def test_paper_meta_falls_back_to_the_filename():
    meta = extract.detect_paper_meta(
        "1. A question with no header.\n", filename="GIT_Gr12_FirstTT (1).pdf")
    assert meta == {"grade": "G12", "subject": "General Information Technology",
                    "title": "First Term Test"}
    assert extract.detect_paper_meta(
        "", filename="G-13, C.maths, Applied , 1st term exam.pdf"
    ) == {"grade": "G13", "subject": "Combined Maths", "title": "1st Term Exam"}


def test_medium_of_instruction_is_not_the_subject():
    meta = extract.detect_paper_meta("Grade 9\nEnglish Medium\nScience\n")
    assert meta["subject"] == "Science"


def test_grade_label_reads_the_way_a_person_writes_it():
    assert extract.grade_label("G12") == "Grade 12"
    assert extract.grade_label("G06") == "Grade 6"
    assert extract.grade_label(None) == ""


def test_detect_paper_meta_rejects_an_impossible_grade():
    meta = extract.detect_paper_meta("Grade 47 monster paper")
    assert meta["grade"] is None


# --------------------------------------------------- extract() on docx

def test_extract_reads_a_docx_end_to_end(tmp_path):
    docx = pytest.importorskip("docx")
    path = tmp_path / "paper.docx"
    d = docx.Document()
    d.add_paragraph("Grade 9 - Mathematics")
    d.add_paragraph("1. Add 12 and 5. (2 marks)")
    d.add_paragraph("2. Subtract 5 from 12. (2 marks)")
    d.save(str(path))

    result = extract.extract(path)
    assert result["file_type"] == "docx"
    assert result["extractor"] == "docx-text"
    assert result["paper"] == {"grade": "G09", "subject": "Maths",
                               "title": None}
    assert [q["question_number"] for q in result["questions"]] == ["1", "2"]
    # Paper-level meta is inherited by each draft.
    assert all(q["grade"] == "G09" and q["subject"] == "Maths"
               for q in result["questions"])


def test_extract_unsupported_type_does_not_raise(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("1. Something. (1 mark)")
    result = extract.extract(path)
    assert result["file_type"] == "other"
    assert result["extractor"].startswith("unsupported")


def test_extract_legacy_doc_asks_for_docx(tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0 legacy binary word file")
    result = extract.extract(path)
    assert "save as .docx" in result["extractor"]
    assert extract.was_incomplete(result["extractor"])


def test_extract_corrupt_pdf_does_not_raise(tmp_path):
    pytest.importorskip("pdfplumber")
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf")
    result = extract.extract(path, ocr=FakeOCR(""))
    assert result["extractor"].startswith("failed")
    assert result["questions"][0]["extraction_confidence"] == "low"


# ------------------------------------------- pictures, scans, equations

class FakeOCR:
    """image -> fixed text, recording what it was handed. Tesseract is
    not needed for these tests - they prove which parts of a file get
    sent to OCR and where the result lands, not how well OCR reads."""

    def __init__(self, text):
        self.text = text
        self.calls = []

    def __call__(self, image):
        self.calls.append(image)
        return self.text


def unavailable_ocr(image):
    raise extract.ExtractionUnavailable("Tesseract is not installed")


def _write_pdf(path, lines, picture=None, pictures=()):
    """A one-page PDF built by hand: Helvetica text at the given heights
    (PDF y, bottom-up) and grey inline pictures at (x, y, width, height)
    - enough to drive each of read_pdf's routes without a PDF library or
    a real paper."""
    ops = [f"BT /F1 12 Tf 50 {y} Td ({s}) Tj ET" for y, s in lines]
    for x, y, w, h in ([picture] if picture else []) + list(pictures):
        ops.append(f"q {w} 0 0 {h} {x} {y} cm BI /W 1 /H 1 /CS /G /BPC 8 ID \x80 EI Q")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def _write_pdf_bold(path, lines):
    """A one-page PDF like _write_pdf(), but each line is (y, text, bold)
    - two font resources, F1 Helvetica and F2 Helvetica-Bold, selected
    per line, so a real bold-font run exists to detect."""
    ops = []
    for y, s, bold in lines:
        font = "/F2" if bold else "/F1"
        ops.append(f"BT {font} 12 Tf 50 {y} Td ({s}) Tj ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R /F2 5 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def test_pdf_bold_option_is_read_as_the_answer_when_no_answer_key_exists(tmp_path):
    """The real Grade 13 Biosystem Technology paper marks its correct
    MCQ option by bolding the whole line - no separate answer key at
    all. Same fallback as the DOCX case, now proven for PDF too."""
    pytest.importorskip("pdfplumber")
    path = tmp_path / "bold.pdf"
    _write_pdf_bold(path, [
        (780, "01.Interpretes weather conditions suitable for", False),
        (760, "biological systems involves,", False),
        (740, "I. Ignoring climate data", False),
        (720, "II. Analyzing conditions for optimal growth", True),
        (700, "III. Focusing only on temperature", False),
        (680, "IV. Avoiding forecasts", False),
        (660, "V. Reducing monitoring", False),
        (640, "02.Water quality determination measures,", False),
        (620, "I. Only temperature", False),
        (600, "II. pH, nutrients and contaminants", True),
        (580, "III. Air pressure", False),
        (560, "IV. Soil texture", False),
        (540, "V. Plant height", False),
    ])
    result = extract.extract(path, ocr=lambda img: "")
    by_number = {q["question_number"]: q for q in result["questions"]}
    assert by_number["01"]["answer"] == "(II)"
    assert by_number["02"]["answer"] == "(II)"


def test_pdf_bold_runs_ignores_text_shorter_than_the_floor(tmp_path):
    """A bold page number or a stray short word is not a paper marking a
    whole MCQ option correct - same length floor as the DOCX case."""
    pytest.importorskip("pdfplumber")
    import pdfplumber
    path = tmp_path / "short.pdf"
    _write_pdf_bold(path, [(780, "Hi", True)])
    with pdfplumber.open(path) as pdf:
        assert extract._pdf_bold_runs(pdf.pages[0]) == []


def test_segment_restarts_numbering_after_a_part_heading():
    text = (
        "Part 1\n1. Define velocity.\n2. Define acceleration.\n"
        "Part 2\nAnswer any 5 questions\n1. Prove the result.\n2. Find the force.\n"
    )
    drafts = extract.segment(text)
    assert [d["question_number"] for d in drafts] == ["1", "2", "1", "2"]
    assert drafts[2]["text"] == "Prove the result."


def test_segment_still_drops_backwards_numbers_without_a_heading():
    text = "1. Choose one.\n(1) red\n(2) blue\n2. Choose again.\n"
    assert [d["question_number"] for d in extract.segment(text)] == ["1", "2"]


def test_page_route_sends_sparse_or_garbled_pages_to_ocr():
    # Only a page number as text - the real model papers' picture pages.
    assert extract._pdf_page_route("9", []) == "page"
    assert extract._pdf_page_route("", [(0, 0, 400, 500)]) == "page"
    assert extract._pdf_page_route("A long enough line (cid:9) here", []) == "page"
    assert extract._pdf_page_route("1. Name the capital of Sri Lanka.", []) == "text"
    assert extract._pdf_page_route(
        "1. Name the capital of Sri Lanka.", [(0, 0, 200, 100)]) == "pictures"


def test_fix_ligatures_restores_the_ti_glyph():
    assert (extract._fix_ligatures("Combined mathema(cid:415)cs ques(cid:415)ons")
            == "Combined mathematics questions")


def test_scanned_pdf_is_rendered_and_ocrd(tmp_path):
    Image = pytest.importorskip("PIL.Image")
    pytest.importorskip("pdfplumber")
    path = tmp_path / "scan.pdf"
    Image.new("L", (400, 500), 255).save(path, "PDF")

    ocr = FakeOCR("Grade 9 Science\n1. Name a planet.\n2. Name a star.")
    result = extract.extract(path, ocr=ocr)

    assert [q["question_number"] for q in result["questions"]] == ["1", "2"]
    assert "ocr (1 of 1 pages)" in result["extractor"]
    # A rendered page image - the old code handed the .pdf file itself
    # to Pillow, which cannot open a PDF.
    assert len(ocr.calls) == 1 and isinstance(ocr.calls[0], Image.Image)
    assert not extract.was_incomplete(result["extractor"])


def test_scanned_pdf_without_tesseract_says_so(tmp_path):
    Image = pytest.importorskip("PIL.Image")
    pytest.importorskip("pdfplumber")
    path = tmp_path / "scan.pdf"
    Image.new("L", (400, 500), 255).save(path, "PDF")

    result = extract.extract(path, ocr=unavailable_ocr)
    assert "1 of 1 pages not read" in result["extractor"]
    assert extract.was_incomplete(result["extractor"])
    assert result["questions"][0]["extraction_confidence"] == "low"


def test_scanned_page_ocr_finds_nothing_is_flagged(tmp_path):
    # Handwriting: Tesseract runs fine and returns nothing. That must not
    # pass for an empty paper.
    Image = pytest.importorskip("PIL.Image")
    pytest.importorskip("pdfplumber")
    path = tmp_path / "handwritten.pdf"
    Image.new("L", (400, 500), 255).save(path, "PDF")

    result = extract.extract(path, ocr=FakeOCR("  "))
    assert "1 of 1 pages not read: OCR found no text" in result["extractor"]
    assert extract.was_incomplete(result["extractor"])


def test_text_pdf_never_touches_ocr(tmp_path):
    pytest.importorskip("pdfplumber")
    path = tmp_path / "typed.pdf"
    _write_pdf(path, [(780, "Grade 9 Science"),
                      (740, "1. Name a planet. (2 marks)"),
                      (700, "2. Name a star. (2 marks)")])
    ocr = FakeOCR("SHOULD NOT APPEAR")
    result = extract.extract(path, ocr=ocr)

    assert ocr.calls == []
    assert result["extractor"] == "pdf-text"
    assert [q["question_number"] for q in result["questions"]] == ["1", "2"]


def test_picture_in_a_typed_pdf_is_read_where_it_sits(tmp_path):
    pytest.importorskip("pdfplumber")
    path = tmp_path / "mixed.pdf"
    # Question 1's text, then a picture below it, then question 2.
    _write_pdf(path, [(780, "Grade 9 Science"),
                      (740, "1. Look at the picture below."),
                      (400, "2. Name a star.")],
               picture=(50, 500, 300, 150))
    ocr = FakeOCR("Which planet is shown here?")
    result = extract.extract(path, ocr=ocr)

    raw = result["raw_text"]
    assert (raw.index("Look at the picture") < raw.index("Which planet")
            < raw.index("Name a star"))
    assert "Which planet is shown here?" in result["questions"][0]["text"]
    assert len(ocr.calls) == 1
    assert "ocr (1 of 1 pages)" in result["extractor"]


def test_a_diagrams_garbled_ocr_is_dropped_not_spliced_in(tmp_path):
    """A real Grade 12 Tamil Physics paper illustrates two questions with
    a velocity-time graph each - small pictures holding no real text, just
    axis labels and lines. Tesseract still "reads" a handful of nonsense
    syllables off them rather than returning nothing, and with no floor on
    a picture's OCR length those syllables spliced themselves into the
    middle of the question's own sentence, as if they were real words.
    A picture producing fewer than _SPARSE_PAGE characters is now treated
    as noise and dropped, the same floor a whole page already uses to
    decide its own text layer is real."""
    pytest.importorskip("pdfplumber")
    path = tmp_path / "graph.pdf"
    _write_pdf(path, [(780, "Grade 12 Physics"),
                      (740, "21. A cart starts from rest at t=0."),
                      (400, "22. The next question.")],
               picture=(50, 500, 300, 150))
    ocr = FakeOCR("l l. l")  # shorter than _SPARSE_PAGE - diagram noise
    result = extract.extract(path, ocr=ocr)

    assert "l l. l" not in result["raw_text"]
    assert len(ocr.calls) == 1  # the picture was still OCR'd, just not kept
    raw = result["raw_text"]
    # Nothing wedged itself between the two real lines - the gap holds
    # only the next line's own printed number, never the dropped picture.
    gap = raw[raw.index("rest at t=0.") + len("rest at t=0."):raw.index("The next question.")]
    assert gap.strip() in ("", "22.")


def test_picture_halves_are_joined_into_one_picture():
    # Coordinates from the real Grade 13 Applied paper's river question.
    top = (121.9, 100.0, 516.5, 211.3)
    bottom = (121.9, 211.3, 516.5, 322.6)
    separate = (121.9, 400.0, 516.5, 500.0)
    assert extract._merge_touching([bottom, separate, top]) == [
        (121.9, 100.0, 516.5, 322.6), separate]


def test_picture_split_by_word_is_ocrd_once_in_a_typed_pdf(tmp_path):
    pytest.importorskip("pdfplumber")
    path = tmp_path / "halves.pdf"
    _write_pdf(path, [(780, "Grade 13 Combined Mathematics"),
                      (740, "1. Answer the question below.")],
               pictures=[(50, 575, 300, 75), (50, 500, 300, 75)])
    ocr = FakeOCR("A river of width a flows with speed u.")
    result = extract.extract(path, ocr=ocr)
    assert len(ocr.calls) == 1
    assert result["raw_text"].count("A river of width") == 1


_MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def _r(s):
    return f"<m:r><m:t>{s}</m:t></m:r>"


def test_omml_fraction_and_root():
    from docx.oxml import parse_xml
    xml = (f"<m:oMath {_MATH_NS}><m:f><m:num>{_r('1')}</m:num><m:den>"
           f"<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/>"
           f"<m:e>{_r('2')}</m:e></m:rad></m:den></m:f></m:oMath>")
    # A bare root counts as one term, the way the real paper reads it:
    # "80 x 1/√2 = 40√2".
    assert extract._omml(parse_xml(xml)) == "1/√2"


def test_omml_subscripts_and_superscripts_under_a_root():
    from docx.oxml import parse_xml

    def v_squared(sub):
        return (f"<m:sSup><m:e><m:sSub><m:e>{_r('v')}</m:e>"
                f"<m:sub>{_r(sub)}</m:sub></m:sSub></m:e>"
                f"<m:sup>{_r('2')}</m:sup></m:sSup>")

    xml = (f"<m:oMath {_MATH_NS}><m:rad><m:deg/><m:e>"
           f"{v_squared('x')}{_r('+')}{v_squared('y')}"
           f"</m:e></m:rad></m:oMath>")
    assert extract._omml(parse_xml(xml)) == "√(v_x^2+v_y^2)"


def test_docx_keeps_equations_inside_the_question(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import parse_xml
    d = docx.Document()
    d.add_paragraph("Grade 12 - Combined Mathematics")
    p = d.add_paragraph("1. Show that 80 x ")
    p._p.append(parse_xml(
        f"<m:oMath {_MATH_NS}><m:f><m:num>{_r('1')}</m:num>"
        f"<m:den>{_r('2')}</m:den></m:f></m:oMath>"))
    p.add_run(" = 40. (5 marks)")
    path = tmp_path / "maths.docx"
    d.save(str(path))

    result = extract.extract(path)
    # "(5 marks)" is read into its own marks field, not left in the text -
    # see test_marks_annotation_is_stripped_from_a_leaf_questions_text.
    assert result["questions"][0]["text"] == "Show that 80 x 1/2 = 40."
    assert result["questions"][0]["marks"] == 5


def test_docx_bold_option_is_read_as_the_answer_when_no_answer_key_exists():
    """Some real papers never print a separate answer key - the correct
    option is simply bolded in place, in the question itself."""
    assert extract._bold_mcq_answer(
        "Which of these is a noble gas?\n(1) Oxygen\n(2) Helium\n"
        "(3) Nitrogen\n(4) Chlorine",
        ["(2) Helium"],  # the whole bolded line, marker included
    ) == "2"
    assert extract._bold_mcq_answer(
        "Which of these is a noble gas?\n(1) Oxygen\n(2) Helium\n"
        "(3) Nitrogen\n(4) Chlorine",
        ["Helium"],  # only the word itself bolded, no marker
    ) == "2"
    # More than one option matches the same short bold word - too
    # ambiguous to guess, so this stays silent rather than picking one.
    assert extract._bold_mcq_answer(
        "Pick the correct pair.\n(1) Iron and Oxygen\n(2) Helium and Neon\n"
        "(3) Iron and Chlorine\n(4) Argon and Xenon",
        ["Iron"],
    ) is None
    # No bold at all, or not an MCQ - nothing to infer.
    assert extract._bold_mcq_answer("Which of these is a noble gas?", []) is None


def test_extract_docx_infers_mcq_answer_from_a_bold_option(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("1. Which of these is a noble gas?")
    d.add_paragraph("(1) Oxygen")
    bold_run = d.add_paragraph().add_run("(2) Helium")
    bold_run.bold = True
    d.add_paragraph("(3) Nitrogen")
    d.add_paragraph("(4) Chlorine")
    d.add_paragraph("2. Which of these is a metal?")
    d.add_paragraph("(1) Oxygen")
    bold_run2 = d.add_paragraph().add_run("(2) Iron")
    bold_run2.bold = True
    d.add_paragraph("(3) Nitrogen")
    d.add_paragraph("(4) Chlorine")
    path = tmp_path / "bold_mcq.docx"
    d.save(str(path))

    result = extract.extract(path)
    by_number = {q["question_number"]: q for q in result["questions"]}
    assert by_number["1"]["answer"] == "(2)"
    assert by_number["2"]["answer"] == "(2)"


def test_docx_roman_numeral_mcq_with_bold_answer_is_not_split(tmp_path):
    """The same real Grade 13 Biosystem Technology paper shape as
    test_mcq_with_verbose_roman_numeral_options_is_not_split_into_parts,
    now as an actual .docx with a bolded correct option - both fixes
    (_OPTION_LINE recognising roman numerals, and bold-run detection)
    have to hold together, not just separately."""
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("01.What is mean by egg's shape index?")
    d.add_paragraph(
        "I. The percentage of ratio between the length and the breadth "
        "that's important in packaging")
    bold_run = d.add_paragraph().add_run(
        "II. The percentage of ratio between the length and the weight "
        "that's important in packaging")
    bold_run.bold = True
    d.add_paragraph(
        "III. The percentage of ratio between the length and the breadth "
        "that's important in deciding the quality while cooking")
    d.add_paragraph(
        "IV. The percentage of ratio between the length and the weight "
        "that's important in deciding the quality while cooking")
    d.add_paragraph(
        "V. The percentage of ration between the breadth and the weight "
        "that's important in deciding the brittleness")
    d.add_paragraph("02.Which of the following is an example for temporary "
                    "controlled environmental structure?")
    d.add_paragraph("I. Web house")
    bold_run2 = d.add_paragraph().add_run("II. Solar greenhouse")
    bold_run2.bold = True
    d.add_paragraph("III. Green house")
    d.add_paragraph("IV. Polythene greenhouse")
    d.add_paragraph("V. Lath house")
    path = tmp_path / "roman_bold.docx"
    d.save(str(path))

    result = extract.extract(path)
    by_number = {q["question_number"]: q for q in result["questions"]}
    assert set(by_number) == {"01", "02"}   # not split into fake parts
    assert by_number["01"]["answer"] == "(II)"
    assert by_number["02"]["answer"] == "(II)"


def test_docx_reads_tables_and_pictures_in_order(tmp_path):
    docx = pytest.importorskip("docx")
    Image = pytest.importorskip("PIL.Image")
    picture = tmp_path / "q.png"
    Image.new("RGB", (200, 60), "white").save(picture)

    d = docx.Document()
    d.add_paragraph("1. Which is the capital?")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A. Colombo"
    table.cell(0, 1).text = "B. Kandy"
    d.add_paragraph("2. Answer the question in the picture.")
    d.add_picture(str(picture))
    path = tmp_path / "paper.docx"
    d.save(str(path))

    ocr = FakeOCR("What is the boiling point of water?")
    result = extract.extract(path, ocr=ocr)
    q1, q2 = result["questions"]
    assert "A. Colombo | B. Kandy" in q1["text"]
    assert "boiling point of water" in q2["text"]
    assert result["extractor"] == "docx-text + ocr (1 image)"


def test_docx_picture_without_tesseract_is_marked_not_dropped(tmp_path):
    docx = pytest.importorskip("docx")
    Image = pytest.importorskip("PIL.Image")
    picture = tmp_path / "q.png"
    Image.new("RGB", (200, 60), "white").save(picture)

    d = docx.Document()
    d.add_paragraph("1. Answer the question in the picture.")
    d.add_picture(str(picture))
    path = tmp_path / "paper.docx"
    d.save(str(path))

    result = extract.extract(path, ocr=unavailable_ocr)
    assert "[image - not read]" in result["questions"][0]["text"]
    assert extract.was_incomplete(result["extractor"])


def test_prepare_for_ocr_flattens_transparency_onto_white():
    Image = pytest.importorskip("PIL.Image")
    # A transparent PNG screenshot: naive greyscale turns it black.
    out = extract._prepare_for_ocr(Image.new("RGBA", (100, 40), (0, 0, 0, 0)))
    assert out.mode == "L"
    assert out.getpixel((0, 0)) == 255
    assert out.width == 400   # upscaled, but capped at 4x


# --------------------------------------------------- check_batch()

class FakeEmbedder:
    """text -> vector, exactly as given. Lets a test decide which drafts
    the model should think are alike."""

    def __init__(self, vectors):
        self.vectors = vectors

    def encode(self, texts, is_query=False):
        return [self.vectors[t] for t in texts]


def test_check_batch_flags_a_near_duplicate_pair():
    drafts = [
        {"id": 1, "text": "Explain how plants make food."},
        {"id": 2, "text": "Describe the way plants produce their own food."},
        {"id": 3, "text": "State Newton's second law of motion."},
    ]
    vectors = {
        drafts[0]["text"]: [1.0, 0.0, 0.0],
        drafts[1]["text"]: [1.0, 0.0, 0.0],   # same meaning per the model
        drafts[2]["text"]: [0.0, 1.0, 0.0],
    }
    out = review.check_batch(drafts, FakeEmbedder(vectors))
    assert out["exact_groups"] == []
    assert len(out["similar_pairs"]) == 1
    pair = out["similar_pairs"][0]
    assert {pair["a"], pair["b"]} == {1, 2}


def test_check_batch_reports_identical_text_as_an_exact_group_only():
    drafts = [
        {"id": 1, "text": "What is 2 + 2?"},
        {"id": 2, "text": "What is 2 + 2?"},
        {"id": 3, "text": "What is the capital of France?"},
    ]
    vectors = {
        "What is 2 + 2?": [1.0, 0.0],
        "What is the capital of France?": [0.0, 1.0],
    }
    out = review.check_batch(drafts, FakeEmbedder(vectors))
    assert out["exact_groups"] == [[1, 2]]
    # The exact pair must not also be listed as a "similar" pair.
    assert out["similar_pairs"] == []


def test_check_batch_catches_the_same_question_under_different_numbers():
    """The real end-to-end case: segment() strips the printed number, so
    a question copy-pasted onto a paper twice at different positions
    hashes identically and lands as an exact group - not lost to the
    numbers guard firing on '3.' vs '7.'."""
    drafts_text = extract.segment(
        "3. Define osmosis. (2 marks)\n7. Define osmosis. (2 marks)\n"
    )
    drafts = [{"id": i + 1, "text": d["text"]} for i, d in enumerate(drafts_text)]
    out = review.check_batch(
        drafts, FakeEmbedder({drafts[0]["text"]: [1.0, 0.0]}),
    )
    assert out["exact_groups"] == [[1, 2]]
    assert out["similar_pairs"] == []


def test_check_batch_needs_at_least_two_drafts():
    out = review.check_batch([{"id": 1, "text": "only one"}], FakeEmbedder({}))
    assert out == {"exact_groups": [], "similar_pairs": []}


def test_check_batch_skips_blank_drafts():
    drafts = [
        {"id": 1, "text": "a real question here"},
        {"id": 2, "text": "   "},
    ]
    out = review.check_batch(drafts, FakeEmbedder({"a real question here": [1.0]}))
    assert out == {"exact_groups": [], "similar_pairs": []}


# ---------------------------------------------------------- storage

def test_document_and_draft_round_trip(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(
            conn, "grade10-science.pdf", file_type="pdf",
            extractor="pdf-text", paper_grade="10", paper_subject="Science",
        )
        store.add_extracted_questions(conn, doc_id, [
            {"seq": 1, "question_number": "1", "text": "Q one", "marks": 2,
             "extraction_confidence": "ok"},
            {"seq": 2, "question_number": "2", "text": "Q two", "marks": None,
             "extraction_confidence": "low"},
        ], fallback_grade="10", fallback_subject="Science")

        drafts = store.intake_questions(conn, doc_id)

    assert [d["seq"] for d in drafts] == [1, 2]
    # Stored once, on the paper, as EdSider's grade - read back as its code.
    assert drafts[0]["grade"] == "G10" and drafts[0]["subject"] == "Science"
    assert drafts[1]["extraction_confidence"] == "low"


def test_update_intake_question_applies_edits(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "raw text", marks=None)
        q_id = store.intake_questions(conn, doc_id)[0]["id"]

        store.update_intake_question(
            conn, q_id, text="corrected text", marks=5, reviewed=1,
        )
        row = store.intake_questions(conn, doc_id)[0]

    assert row["text"] == "corrected text"
    assert row["marks"] == 5
    assert row["reviewed"] == 1


def test_teacher_reviewed_is_separate_from_the_data_entry_tick(local_db):
    """Two roles, two ticks - a teacher confirming a question must not
    silently mark it reviewed for data-entry, or the reverse."""
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "raw text")
        q_id = store.intake_questions(conn, doc_id)[0]["id"]

        store.update_intake_question(conn, q_id, reviewed=1)
        row = store.intake_questions(conn, doc_id)[0]
        assert row["reviewed"] == 1
        assert row["teacher_reviewed"] == 0

        store.update_intake_question(conn, q_id, teacher_reviewed=1)
        row = store.intake_questions(conn, doc_id)[0]
        assert row["reviewed"] == 1
        assert row["teacher_reviewed"] == 1


def test_question_comments_round_trip_per_document(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Q one")
        store.add_intake_question(conn, doc_id, 2, "Q two")
        q1, q2 = store.intake_questions(conn, doc_id)

        store.add_question_comment(conn, q1["id"], "Ms. Perera",
                                    "Option 3 should be marked correct.")
        store.add_question_comment(conn, q1["id"], "Ms. Perera",
                                    "Actually option 2 is right.")

        by_question = store.question_comments_for_document(conn, doc_id)

    assert [c["comment"] for c in by_question[q1["id"]]] == [
        "Option 3 should be marked correct.",
        "Actually option 2 is right.",
    ]
    assert by_question[q1["id"]][0]["author"] == "Ms. Perera"
    assert q2["id"] not in by_question


def test_question_images_round_trip_per_document(local_db):
    """A reviewer attaching a picture of the actual diagram/option/table
    to a question, its answer, or one MCQ option - for content
    extraction could not turn into meaningful text on its own."""
    fake_png = b"\x89PNG\r\n\x1a\n" + b"fake image bytes"
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Name the cut flower shown.")
        store.add_intake_question(conn, doc_id, 2, "Q two")
        q1, q2 = store.intake_questions(conn, doc_id)

        img_id = store.add_question_image(
            conn, q1["id"], "question", fake_png, content_type="image/png")
        store.add_question_image(
            conn, q1["id"], "option:I", b"option I picture bytes")

        by_question = store.question_images_for_document(conn, doc_id)
        blob = store.question_image_blob(conn, img_id)

    assert [i["slot"] for i in by_question[q1["id"]]] == ["question", "option:I"]
    assert by_question[q1["id"]][0]["content_type"] == "image/png"
    assert q2["id"] not in by_question
    assert blob == fake_png


def test_delete_question_image_removes_it(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Q one")
        q1 = store.intake_questions(conn, doc_id)[0]
        img_id = store.add_question_image(conn, q1["id"], "answer", b"bytes")

        store.delete_question_image(conn, img_id)
        by_question = store.question_images_for_document(conn, doc_id)
        blob = store.question_image_blob(conn, img_id)

    assert by_question == {}
    assert blob is None


def test_upsert_teacher_records_first_visit_and_keeps_it_on_later_ones(local_db):
    """A teacher's first visit to their review page records them on their
    staff profile (and so in the assignment dropdown); later visits only
    refresh last_seen_at."""
    with store.connect() as conn:
        store.upsert_teacher(conn, "dilhani@edsider.lk", "Dilhani Data")
        first_seen = store.lookup_staff(conn, "dilhani@edsider.lk")["first_seen_at"]

        store.upsert_teacher(conn, "dilhani@edsider.lk", "Dilhani D.")
        store.upsert_teacher(conn, "nimal@edsider.lk", "Nimal Teacher")
        dilhani = store.lookup_staff(conn, "dilhani@edsider.lk")
        assignable = {t["identity"] for t in store.assignable_teachers(conn)}

    assert assignable == {"dilhani@edsider.lk", "nimal@edsider.lk"}
    # A stored name is kept - EdSider's or a superadmin's always wins over
    # what a sign-in reports (see store.upsert_teacher) - and first_seen_at
    # does not move.
    assert dilhani["name"] == "Dilhani Data"
    assert dilhani["first_seen_at"] == first_seen
    assert dilhani["last_seen_at"] >= first_seen


def test_assign_document_moves_it_to_teacher_review(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf", extractor="pdf-text")
        store.update_intake_document(conn, doc_id, status="checked")

        store.assign_document(conn, doc_id, "dilhani@edsider.lk")
        doc = store.get_intake_document(conn, doc_id)
        mine = store.documents_assigned_to(conn, "dilhani@edsider.lk")
        someone_elses = store.documents_assigned_to(conn, "nimal@edsider.lk")

    assert doc["status"] == "teacher_review"
    assert doc["assigned_to"] == "dilhani@edsider.lk"
    assert [d["id"] for d in mine] == [doc_id]
    assert someone_elses == []


def test_promote_document_copies_questions_into_the_real_bank(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(
            conn, "grade10-git.pdf", uploaded_by="data.entry@edsider.lk")
        store.add_intake_question(
            conn, doc_id, 1, "A coastal town installs a flood warning system.",
            grade="G12", subject="General Information Technology",
        )
        store.add_intake_question(
            conn, doc_id, 2, "Identify the components.", marks=3,
            answer="Rainfall sensors, timestamps.",
            parent_seq=1, depth=1,
        )
        store.update_intake_document(conn, doc_id, status="review_done")

        promoted = store.promote_document(conn, doc_id)
        doc = store.get_intake_document(conn, doc_id)
        parent_id = f"intake-{doc_id}-1"
        child_id = f"intake-{doc_id}-2"
        parent = conn.execute(
            "SELECT * FROM questions WHERE id = ?", (parent_id,)).fetchone()
        child = conn.execute(
            "SELECT * FROM questions WHERE id = ?", (child_id,)).fetchone()

    assert promoted == 2
    assert doc["status"] == "published"
    assert parent["shape"] == "parent"
    assert parent["children_raw"] == child_id
    assert child["shape"] == "sub-question"
    assert child["parent_id"] == parent_id
    assert child["marks"] == 3
    assert child["answer"] == "Rainfall sensors, timestamps."
    assert child["contributor"] == "data.entry@edsider.lk"
    assert child["is_answer_blob"] == 0


def test_promote_document_refuses_before_review_is_done(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Q one")
        store.update_intake_document(conn, doc_id, status="teacher_review")

        with pytest.raises(ValueError, match="review is done"):
            store.promote_document(conn, doc_id)


def test_update_intake_question_rejects_unknown_columns(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "t")
        q_id = store.intake_questions(conn, doc_id)[0]["id"]
        with pytest.raises(ValueError, match="editable"):
            store.update_intake_question(conn, q_id, document_id=999)


def test_delete_intake_question_drops_a_spurious_row(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "a heading read as a question")
        store.add_intake_question(conn, doc_id, 2, "a real question")
        q_id = store.intake_questions(conn, doc_id)[0]["id"]

        store.delete_intake_question(conn, q_id)
        left = store.intake_questions(conn, doc_id)

    assert [d["seq"] for d in left] == [2]


def test_raw_text_is_kept_so_a_paper_can_be_resplit(local_db):
    """Without this a reviewer has to re-upload the file just to pick up
    an extractor fix - which is exactly how two real papers ended up
    stuck on a stale, wrong segmentation."""
    with store.connect() as conn:
        doc_id = store.create_intake_document(
            conn, "paper.docx", raw_text="1. Alpha.\n2. Beta.\n",
            split_style="number",
        )
        doc = store.get_intake_document(conn, doc_id)
    assert doc["raw_text"] == "1. Alpha.\n2. Beta.\n"
    assert doc["split_style"] == "number"


def test_replace_intake_questions_swaps_the_whole_set(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.docx")
        store.add_intake_question(conn, doc_id, 1, "one big wrong blob")
        assert len(store.intake_questions(conn, doc_id)) == 1

        fresh = extract.segment("1. Alpha question.\n2. Beta question.\n")
        store.replace_intake_questions(conn, doc_id, fresh,
                                       fallback_grade="13")
        rows = store.intake_questions(conn, doc_id)

    assert [r["question_number"] for r in rows] == ["1", "2"]
    assert rows[0]["text"] == "Alpha question."
    assert rows[0]["grade"] == "13"


def test_parts_and_answers_round_trip(local_db):
    parsed = extract.structure(_TWO_PART_PAPER)
    with store.connect() as conn:
        doc_id = store.create_intake_document(
            conn, "GIT_Gr12_FirstTT.pdf", paper_title="First Term Examination",
            sections=parsed["sections"],
        )
        store.add_extracted_questions(conn, doc_id, parsed["questions"])
        doc = store.get_intake_document(conn, doc_id)
        rows = store.intake_questions(conn, doc_id)
        # A part is a row its questions point at: one still holding
        # questions stays when a list leaves it out; an empty one goes.
        store.update_intake_document(conn, doc_id,
                                     sections=parsed["sections"] + [{"label": "Part III"}])
        store.update_intake_document(conn, doc_id,
                                     sections=parsed["sections"][:1])
        after = store.intake_sections(store.get_intake_document(conn, doc_id))

    assert doc["paper_title"] == "First Term Examination"
    assert [s["label"] for s in store.intake_sections(doc)] == [
        "Part I", "Part II"]
    assert [r["section"] for r in rows[:2]] == ["Part I", "Part I"]
    assert {r["section"] for r in rows[2:]} == {"Part II"}
    assert rows[0]["answer"] == "(2)"
    # A question's parts survive the round trip, still pointing at it.
    q01 = next(r for r in rows if r["question_number"] == "01")
    under = [r for r in rows if r["parent_seq"] == q01["seq"]]
    assert [r["question_number"] for r in under] == ["01(1)", "01(2)"]
    assert {r["depth"] for r in under} == {1}
    assert [s["label"] for s in after] == ["Part I", "Part II"]


def test_deleting_a_question_removes_its_parts(local_db):
    """A part left behind when its question goes is an orphan the review
    page cannot place under anything."""
    parsed = extract.structure(_TWO_PART_PAPER)
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "GIT.pdf")
        store.add_extracted_questions(conn, doc_id, parsed["questions"])
        rows = store.intake_questions(conn, doc_id)
        q01 = next(r for r in rows if r["question_number"] == "01")
        store.delete_intake_question(conn, q01["id"])
        left = [r["question_number"] for r in store.intake_questions(conn, doc_id)]

    assert not any(n and n.startswith("01(") for n in left)
    assert left == ["1", "2", "02"]


def test_deleting_a_question_with_comments_and_images_does_not_fail(local_db):
    """intake_question_comments and intake_question_images reference the
    question they are attached to - deleting a question that has either
    must clean them up first, not hit a foreign-key violation."""
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Q one")
        q1 = store.intake_questions(conn, doc_id)[0]
        store.add_question_comment(conn, q1["id"], "Ms. Perera", "Wrong answer.")
        store.add_question_image(conn, q1["id"], "question", b"bytes")

        store.delete_intake_question(conn, q1["id"])

        left = store.intake_questions(conn, doc_id)
        comments = store.question_comments_for_document(conn, doc_id)
        images = store.question_images_for_document(conn, doc_id)

    assert left == []
    assert comments == {}
    assert images == {}


def test_replacing_questions_with_comments_and_images_does_not_fail(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.add_intake_question(conn, doc_id, 1, "Q one")
        q1 = store.intake_questions(conn, doc_id)[0]
        store.add_question_comment(conn, q1["id"], "Ms. Perera", "Wrong answer.")
        store.add_question_image(conn, q1["id"], "question", b"bytes")

        store.replace_intake_questions(
            conn, doc_id, [{"seq": 1, "question_number": "1", "text": "Q new"}])

        left = [r["text"] for r in store.intake_questions(conn, doc_id)]
        comments = store.question_comments_for_document(conn, doc_id)

    assert left == ["Q new"]
    assert comments == {}


def test_a_paper_read_before_parts_existed_has_no_parts(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "old.pdf")
        assert store.intake_sections(
            store.get_intake_document(conn, doc_id)) == []


def test_update_intake_document_sets_paper_level_fields(local_db):
    with store.connect() as conn:
        doc_id = store.create_intake_document(conn, "p.pdf")
        store.update_intake_document(
            conn, doc_id, paper_grade="8", paper_subject="History",
            status="checked",
        )
        doc = store.get_intake_document(conn, doc_id)

    # Stored as EdSider's grade id, read back as EdSider's code.
    assert doc["paper_grade"] == "G08"
    assert doc["paper_subject"] == "History"
    assert doc["status"] == "checked"
